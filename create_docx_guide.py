import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout_box(doc, title, text, bg_hex="F0F4F8", border_hex="1E3A8A"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"💡 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10.5)
    run_b.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph()

def build_interview_doc():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("MOODYTRIPS AI — FRESHER INTERVIEW PREPARATION GUIDE")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy Blue
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Comprehensive Technical Q&A, Multi-Agent Architectures, LangGraph vs LangChain, Database Checkpoints & Interview Strategies\nCandidate Focus: B.Tech (Mathematics & Computing) Freshers Interview")
    sub_run.font.size = Pt(12)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    add_callout_box(
        doc,
        "How to Use This Document for Interviews",
        "This guide is specifically tailored for freshers defending MoodyTrips AI in technical interviews. It covers everything from 30-second pitches to deep technical trade-offs between LangChain & LangGraph, PostgreSQL checkpoint state persistence, custom API tools, and future scalability improvements. Review Section 1 before standard HR/Technical rounds and Sections 2-7 for system design and coding discussions.",
        bg_hex="EFF6FF",
        border_hex="2563EB"
    )
    
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88) # Teal
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        return h

    def add_p(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
        p.add_run(text)
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # SECTION 1: EXECUTIVE ELEVATOR PITCH & RESUME FRAMING
    # ==========================================
    add_h1("SECTION 1: Executive Elevator Pitch & Resume Framing")
    
    add_h2("1.1 The 30-Second Elevator Pitch")
    add_p(
        "\"MoodyTrips AI is an intelligent, multi-agent travel planning platform built using LangGraph, LangChain, FastAPI, and PostgreSQL. Unlike traditional monolithic chatbots that struggle with context and rigid booking rules, MoodyTrips decouples travel planning into 7 autonomous specialized agents—handling flights, hotels, live Indian railway schedules, weather forecasts, and budget optimizations. The system maintains full conversational state using PostgreSQL checkpointing and serves a real-time responsive web UI powered by Groq's high-speed Llama-3.3-70B model.\""
    )
    
    add_h2("1.2 The 2-Minute In-Depth Project Breakdown")
    add_p(
        "When an interviewer asks: 'Walk me through your MoodyTrips project', hit these 4 structured milestones:"
    )
    add_bullet(" Built a multi-agent state graph using LangGraph where each node represents a domain-specific agent (Flight Agent, Hotel Agent, Ground Transport Agent, Weather Agent, Budget Optimizer Agent, Itinerary Agent, Final Response Agent).", "1. Multi-Agent Architecture:")
    add_bullet(" Integrated real-time external APIs: AviationStack for flight route detection, Tavily Search for live hotel/weather research, and RailKit/IRCTC for live Indian train schedules and seat availability.", "2. External Tool Integration:")
    add_bullet(" Engineered state persistence with PostgreSQL PostgresSaver checkpointer, allowing users to pause, resume, and manage multiple distinct travel planning chat threads seamlessly across sessions with JWT authentication.", "3. Persistence & Auth:")
    add_bullet(" Built a asynchronous FastAPI backend coupled with Jinja2/vanilla JavaScript dynamic UI and secured endpoints using Bcrypt password hashing and PyJWT tokens.", "4. Full-Stack Delivery:")

    add_h2("1.3 Resume Standard Bullet Points")
    add_bullet("Architected a multi-agent AI travel assistant using LangGraph & LangChain, orchestrating 7 specialized agents to synthesize personalized flight, train, hotel, and weather itineraries.")
    add_bullet("Integrated real-time APIs (AviationStack, RailKit, Tavily) and implemented custom Regex extraction for Indian Railways seat availability & dynamic date routing.")
    add_bullet("Implemented thread-level session state persistence using PostgreSQL PostgresSaver checkpointer, enabling cross-session chat resumption and multi-tenant thread management.")
    add_bullet("Engineered a production-ready FastAPI web backend with JWT-based user authentication, Bcrypt password hashing, and clean RESTful endpoints.")
    add_bullet("Leveraged Groq LLM (Llama-3.3-70B-Versatile) for low-latency reasoning and dynamic budget optimization tables.")

    # ==========================================
    # SECTION 2: SYSTEM ARCHITECTURE & DIAGRAMS
    # ==========================================
    add_h1("SECTION 2: System Architecture & Diagrams")
    
    add_h2("2.1 High-Level End-to-End Architecture")
    add_p("The end-to-end flow of data from the User UI down to the LLM and DB checkpointer:")
    
    add_code_block(
"""+-------------------------------------------------------------------------+
|                              USER BROWSER                               |
|        (HTML5 / CSS3 / Vanilla JS / Dynamic Markdown Rendering)         |
+-------------------------------------------------------------------------+
                                    |
                            HTTP POST /api/travel
                        (Bearer JWT Auth Header)
                                    v
+-------------------------------------------------------------------------+
|                             FASTAPI BACKEND                             |
|  - Auth Middleware (JWT Decode & Bcrypt verification)                   |
|  - Chat Thread Management (database.py -> user_chats, chat_messages)     |
|  - Invokes LangGraph StateGraph Execution                               |
+-------------------------------------------------------------------------+
                                    |
                                    v
+-------------------------------------------------------------------------+
|                        LANGGRAPH WORKFLOW ENGINE                        |
|                                                                         |
|  START -> [Flight Agent] -> [Hotel Agent] -> [Ground Transport Agent]   |
|            (AviationStack)    (Tavily API)        (RailKit IRCTC)       |
|                                                          |              |
|  END <- [Final Agent] <- [Itinerary Agent] <- [Budget Agent] <- [Weather]|
|                                                 (Groq Llama-3.3)        |
+-------------------------------------------------------------------------+
             |                                              |
             v                                              v
+--------------------------+                  +---------------------------+
|   POSTGRES CHECKPOINTER  |                  |      GROQ LLM API         |
| (PostgresSaver / psycopg)|                  | (llama-3.3-70b-versatile) |
+--------------------------+                  +---------------------------+)"""
    )
    
    add_h2("2.2 LangGraph State Machine Graph")
    add_p("The state graph executes strictly in a structured sequential workflow, accumulating state at each step:")
    
    add_code_block(
"""[START]
   │
   ▼
[flight_agent] ----------> Searches AviationStack / extracts flight options
   │
   ▼
[hotel_agent] -----------> Searches Tavily API for accommodations / hostels
   │
   ▼
[ground_transport_agent]-> Queries RailKit for Indian train seat availability & fares
   │
   ▼
[weather_agent] ---------> Searches Tavily for destination weather & season data
   │
   ▼
[budget_agent] ----------> Invokes Groq LLM to compute markdown cost breakdown table
   │
   ▼
[itinerary_agent] -------> Invokes Groq LLM to construct day-wise travel schedule
   │
   ▼
[final_agent] -----------> Synthesizes all state variables into final Markdown response
   │
   ▼
 [END]"""
    )

    add_h2("2.3 State Schema (TravelState)")
    add_p("The `TravelState` TypedDict defines the shared memory passed across every agent node:")
    
    add_code_block(
"""class TravelState(TypedDict):
    messages: Annotated[list[AnyMessage], operator.add] # State reducer appends messages
    user_query: str                  # Original natural language request
    flight_results: str              # Aggregated flight details
    hotel_results: str               # Aggregated hotel/hostel details
    ground_transport_results: str    # Train/Bus live schedules & seat info
    weather_results: str             # Climate and weather forecast data
    budget_breakdown: str            # Calculated itemized cost table
    itinerary: str                   # Day-wise plan generated by LLM
    llm_calls: int                   # Telemetry counter tracking total LLM invocations"""
    )

    # ==========================================
    # SECTION 3: TECH STACK DEEP DIVE
    # ==========================================
    add_h1("SECTION 3: Technology Stack Deep Dive")
    
    # Table for Tech Stack
    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    headers = ["Layer", "Technology Used", "Role in MoodyTrips AI"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    data = [
        ("Orchestration", "LangGraph & LangChain Core", "Stateful multi-agent workflow, node graphs, state reducers, Postgres checkpointer."),
        ("LLM Engine", "Groq (llama-3.3-70b-versatile)", "Ultra-low latency inference for budget optimization, itinerary creation, and response synthesis."),
        ("Web Framework", "FastAPI (Python 3.10+)", "Asynchronous HTTP backend, CORS, JWT security middleware, static file serving, Jinja2 template rendering."),
        ("Database", "PostgreSQL (psycopg 3)", "Stores user accounts, thread histories, chat logs, and LangGraph checkpointer snapshots."),
        ("APIs & Tools", "AviationStack, RailKit, Tavily", "Fetches real-time flight routes, Indian Railways live seat availability/fares, and web lodging search.")
    ]
    
    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_hex = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text
            set_cell_background(row_cells[col_idx], bg_hex)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # SECTION 4: LANGGRAPH VS LANGCHAIN (CRUCIAL DEEP DIVE)
    # ==========================================
    add_h1("SECTION 4: Deep Dive — LangGraph vs LangChain")
    
    add_h2("4.1 Why LangGraph over plain LangChain?")
    add_p(
        "This is one of the MOST IMPORTANT questions in modern AI engineering interviews. Be ready to give a sharp, confident answer:"
    )
    
    add_bullet("LangChain Chains (`LLMChain`, `SequentialChain`) are Directed Acyclic Graphs (DAGs). They execute linearly step-by-step. They lack native support for cyclic loops (e.g. self-correction, retry loops, multi-turn state transitions).", "LangChain Chains are Linear / Monolithic: ")
    add_bullet("LangGraph models workflows as State Machines with explicit nodes (functions), edges (control flow), and a central state schema (`TypedDict`). It allows loops, conditional branching, human-in-the-loop validation, and parallel execution.", "LangGraph introduces Cyclic Graphs & State Machines: ")
    add_bullet("Plain LangChain agents depend on a single `AgentExecutor` loop which passes full tool histories back and forth. LangGraph allows specialized micro-agents where each node has its own isolated prompt, tools, and responsibilities, keeping token usage low and precision high.", "Decoupled Specialized Multi-Agents: ")
    add_bullet("LangGraph features built-in state persistence (`PostgresSaver`, `MemorySaver`). Every step execution automatically checkpoints the graph state into PostgreSQL keyed by `thread_id`.", "First-Class Checkpointing: ")

    add_h2("4.2 LangChain vs LangGraph Comparison Matrix")
    
    table2 = doc.add_table(rows=7, cols=3)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    
    headers2 = ["Feature", "Standard LangChain", "LangGraph (Used in MoodyTrips)"]
    hdr_cells2 = table2.rows[0].cells
    for i, title in enumerate(headers2):
        hdr_cells2[i].text = title
        set_cell_background(hdr_cells2[i], "0D9488") # Teal header
        p = hdr_cells2[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    matrix_data = [
        ("Architecture Paradigm", "DAG (Linear Chain)", "State Machine Graph (Cyclic / Branching)"),
        ("State Management", "Implicit pass-through of variables", "Explicit central `TypedDict` state schema"),
        ("State Reducers", "Overwrites previous state variables", "Custom reducers like `Annotated[list, operator.add]`"),
        ("Multi-Agent Coordination", "Difficult; requires nested AgentExecutors", "Native; nodes are independent specialized agents"),
        ("Cycles & Loops", "Requires hacky recursion", "Native graph cycles (e.g. agent feedback loops)"),
        ("State Checkpointing", "Requires manual custom DB wrapper", "Built-in `PostgresSaver` per `thread_id` snapshot")
    ]
    
    for row_idx, row_data in enumerate(matrix_data, start=1):
        row_cells = table2.rows[row_idx].cells
        bg_hex = "F4FBF7" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text
            set_cell_background(row_cells[col_idx], bg_hex)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # SECTION 5: AGENT-BY-AGENT IMPLEMENTATION BREAKDOWN
    # ==========================================
    add_h1("SECTION 5: Agent-by-Agent Detailed Implementation")
    
    add_h2("5.1 Flight Agent (`flight_agent`)")
    add_bullet("Parses user destination/origin mentions using regex.", "Functionality: ")
    add_bullet("Calls `search_flights(query)` in `tools/flight_tool.py` using AviationStack API.", "Tool Integration: ")
    add_bullet("Updates `flight_results` in `TravelState` and increments `llm_calls` telemetry counter.", "State Mutation: ")

    add_h2("5.2 Hotel Agent (`hotel_agent`)")
    add_bullet("Constructs structured search queries e.g. 'Best hotels, hostels, and homestays for [destination]'.", "Functionality: ")
    add_bullet("Calls Tavily Web Search API to retrieve curated lodging recommendations.", "Tool Integration: ")
    add_bullet("Populates `hotel_results` state field.", "State Mutation: ")

    add_h2("5.3 Ground Transport Agent (`ground_transport_agent`)")
    add_bullet("Extracts travel dates (e.g. '15 Aug' or '15-08-2026') using regex matching.", "Regex Location/Date Extraction: ")
    add_bullet("Extracts Origin and Destination from query. Defaults to Delhi -> Jaipur if unspecified.", "Station Extraction: ")
    add_bullet("Calls `search_trains_between_stations()` and queries real-time seat availability (`get_seat_availability()`) across 3A, CC, SL classes via RailKit API.", "IRCTC Integration: ")
    add_bullet("If no direct trains exist, falls back gracefully to Tavily search for bus and driving directions.", "Fallback Handling: ")

    add_h2("5.4 Weather Agent (`weather_agent`)")
    add_bullet("Executes climate and weather query search via Tavily API for the requested location/dates.", "Functionality: ")
    add_bullet("Populates `weather_results` in state for downstream packing advice.", "State Mutation: ")

    add_h2("5.5 Budget Optimizer Agent (`budget_agent`)")
    add_bullet("Acts as a value-focused travel budget planner.", "Role: ")
    add_bullet("Consolidates flight, train, and hotel costs. Analyzes cheaper alternatives (e.g., recommend trains over flights for domestic travel, Zostel/hostels over hotels).", "Logic: ")
    add_bullet("Invokes Groq LLM to generate an itemized Markdown table displaying Category, Estimated Cost, and Optimization Strategy.", "Output: ")

    add_h2("5.6 Itinerary Agent (`itinerary_agent`)")
    add_bullet("Generates a comprehensive day-by-day travel plan balancing sightseeing, dining, transit, and weather forecast.", "Role: ")
    add_bullet("Applies travel dates extracted by previous agents directly into day numbers.", "Date Alignment: ")

    add_h2("5.7 Final Response Agent (`final_agent`)")
    add_bullet("Synthesizes all state variables (`flight_results`, `ground_transport_results`, `hotel_results`, `weather_results`, `budget_breakdown`, `itinerary`).", "Role: ")
    add_bullet("Formats response into 8 clean structured sections including booking instructions for Skyscanner, IRCTC, and hostels.", "Formatting: ")

    # ==========================================
    # SECTION 6: TOP 25 FRESHER INTERVIEW QUESTIONS & ANSWERS
    # ==========================================
    add_h1("SECTION 6: Top 25 Fresher Interview Questions & Answers")
    
    qa_list = [
        # Category A: Project Overview
        ("Q1: What inspired you to build MoodyTrips AI?", 
         "Traditional travel platforms force users to jump between multiple tabs—checking Skyscanner for flights, IRCTC for Indian trains, Booking.com for hotels, and Google for weather. Generic LLM chatbots often hallucinate prices or lack real-time data. I built MoodyTrips AI to solve this by creating a multi-agent system where specialized agents interface with real-time APIs and collaborate via a state graph to deliver a single, accurate, actionable travel plan."),
        
        ("Q2: Why did you choose a Multi-Agent architecture instead of one single LLM prompt?", 
         "Single-prompt LLM systems suffer from 'prompt stuffing' and high context degradation. Asking a single prompt to search flights, calculate train availability, check weather, optimize budgets, and format itineraries results in missed tools or hallucinations. By decoupling into 7 specialized agents, each agent performs one task deterministically with isolated context, resulting in higher accuracy, cleaner modular code, and cheaper token usage."),

        ("Q3: What role does Groq LLM play in this project?", 
         "Groq provides ultra-fast LLM inference hardware (LPU - Language Processing Unit). We use `llama-3.3-70b-versatile` via Groq. Its high output speed (over 250 tokens/sec) is critical for multi-agent applications where multiple LLM calls occur per user request, reducing total latency from 15+ seconds down to under 3 seconds."),

        ("Q4: How does MoodyTrips handle user authentication and session management?", 
         "We implemented JWT (JSON Web Tokens) with HS256 algorithm and bcrypt password hashing. When users register/login via `/api/auth/register` or `/api/auth/login`, FastAPI generates a signed JWT. The frontend attaches this token in the `Authorization: Bearer <token>` header for subsequent requests. FastApi validates the token, extracts `user_id`, and queries user-specific chat threads."),

        ("Q5: What is the database schema for MoodyTrips?", 
         "MoodyTrips uses PostgreSQL with 3 primary core tables:\n1. `users`: Stores `id`, `email` (unique), `password_hash` (bcrypt), `created_at`.\n2. `user_chats`: Stores `id`, `user_id` (FK to users), `thread_id` (unique string), `title`, `created_at`.\n3. `chat_messages`: Stores `id`, `thread_id` (FK to user_chats), `role` ('user'/'assistant'), `content`, `created_at`.\nIn addition, LangGraph automatically creates its own checkpoint tables (`checkpoints`, `checkpoint_blobs`, `checkpoint_writes`) via `PostgresSaver`."),

        # Category B: LangGraph & AI Concepts
        ("Q6: Explain what `StateGraph` and `TravelState` mean in LangGraph.", 
         "`StateGraph` is the core graph compiler in LangGraph. It is initialized with a state structure (`TravelState`), which is a Python `TypedDict`. `TravelState` acts as the shared central state passed from node to node. Nodes are standard Python functions that receive the current state, perform operations, and return a dictionary of state keys to update."),

        ("Q7: What are Reducers in LangGraph state, and why did you use `operator.add`?", 
         "By default, when a node returns a dictionary key (e.g. `hotel_results: '...'`), LangGraph overwrites the existing value. A Reducer dictates how updates are merged. In `TravelState`, `messages` is defined as `Annotated[list[AnyMessage], operator.add]`. The reducer `operator.add` tells LangGraph to APPEND new messages to the existing list rather than overwriting it, preserving full chat history."),

        ("Q8: How does PostgreSQL Checkpointing (`PostgresSaver`) work in LangGraph?", 
         "`PostgresSaver` is a persistence checkpointer. After every agent node finishes execution, `PostgresSaver` serializes the current graph state snapshot into PostgreSQL indexed by `thread_id`. If the system restarts or the user returns days later, LangGraph loads the graph state directly from PostgreSQL using the `thread_id` config parameter."),

        ("Q9: What happens if an external API (like AviationStack or RailKit) fails or times out?", 
         "We implemented graceful fallback degradation. For instance, in `ground_transport_agent`, if the RailKit IRCTC API fails or finds no direct trains, the agent catches the condition and automatically falls back to invoking the Tavily web search tool to retrieve bus routes, driving directions, or alternative transport schedules, ensuring the pipeline never crashes."),

        ("Q10: What is the difference between synchronous and asynchronous nodes in LangGraph?", 
         "Synchronous nodes execute sequentially on the main thread, blocking until completed. Asynchronous nodes (`async def node_name(state)`) run non-blocking on asyncio event loops. While our current nodes run sequentially in Python, LangGraph natively supports async node execution for parallel fan-out operations."),

        # Category C: Backend & Engineering
        ("Q11: Why did you choose FastAPI over Flask or Django?", 
         "FastAPI was chosen because:\n1. Asynchronous Support: Built on Starlette and ASGI, making it extremely fast for handling concurrent async I/O requests.\n2. Automatic Data Validation: Uses Pydantic for request payload parsing and schema validation.\n3. Automatic API Docs: Generates interactive Swagger UI (`/docs`) out of the box.\n4. Lightweight: Unlike Django, FastAPI does not impose heavy ORM boilerplate when paired with `psycopg 3`."),

        ("Q12: How do you handle password security in MoodyTrips?", 
         "Passwords are NEVER stored in plain text. When a user registers, we use `bcrypt.hashpw()` with a salt to hash the password before inserting into the `users` table. During login, `bcrypt.checkpw()` compares the plain text input against the stored hash safely against timing attacks."),

        ("Q13: How does regex help in your Ground Transport Agent?", 
         "Users type natural queries like 'Plan a trip from Delhi to Manali on 15th Aug'. Standard string splitting fails here. We use regex (`re.search`) with date patterns (`r'(\\d{1,2}(?:st|nd|rd|th)?\\s+(?:Jan|Feb|Mar|...))'`) and location mention parsers (`find_location_mentions`) to dynamically isolate dates and origin/destination cities to query IRCTC APIs cleanly."),

        ("Q14: Explain `psycopg 3` and why `dict_row` was used.", 
         "`psycopg 3` is the modern Python PostgreSQL driver. We used `row_factory=dict_row` when opening connections so database queries return native Python dictionaries (`{'id': 1, 'email': '...'}`) instead of tuple indices (`(1, '...')`), making codebase data access readable and bug-free."),

        ("Q15: How is thread creation managed when a user sends a message?", 
         "In `app.py`, when a POST request hits `/api/travel`, FastAPI extracts `thread_id`. If `thread_id` is missing or does not belong to the authenticated `user_id`, FastAPI generates a new UUID thread ID (`user_<uuid>`), extracts destination location mentions to generate a title (e.g. 'Trip to Bali'), inserts a row into `user_chats`, and passes `thread_id` to LangGraph."),

        # Category D: Advanced & Interview Scenarios
        ("Q16: How would you parallelize agent execution in MoodyTrips?", 
         "Currently, nodes run sequentially: Flight -> Hotel -> Ground Transport -> Weather. However, Flight, Hotel, Ground Transport, and Weather agents are completely independent! In LangGraph, we can connect `START` directly to all 4 agents in parallel edges. LangGraph will execute them concurrently using `asyncio.gather()`, reducing total execution time by ~70%!"),

        ("Q17: How would you implement Human-in-the-Loop in LangGraph?", 
         "LangGraph has native support for `interrupt_before` or `interrupt_after`. For example, after `budget_agent` generates a cost breakdown, we can interrupt execution, present the budget options to the user on UI, get their approval/modification, and resume the graph execution from that exact checkpoint state."),

        ("Q18: What is prompt engineering, and how did you structure prompts for your agents?", 
         "Prompt engineering involves crafting clear system instructions, constraints, and output formatting rules for LLMs. In MoodyTrips, each agent receives explicit System Messages defining its persona (e.g. 'You are an expert value-oriented travel budget planner') and Human Messages containing isolated context variables (e.g., flight and hotel state outputs) to ensure deterministic markdown tables."),

        ("Q19: How do you prevent SQL Injection and CORS security issues?", 
         "1. SQL Injection: All SQL queries in `database.py` use parameterized queries with `%s` placeholders handled by `psycopg 3`, preventing malicious SQL insertion.\n2. CORS: FastAPI supports `CORSMiddleware` to restrict API access strictly to allowed origins."),

        ("Q20: How does Tavily Search API differ from standard Google Search scrapers?", 
         "Tavily is a search engine built specifically for AI agents and RAG (Retrieval-Augmented Generation). Unlike Google scrapers that return messy raw HTML, Tavily cleans HTML, removes boilerplate ads/navigation, extracts relevant context snippets, and returns structured JSON optimal for LLM prompts."),

        # Category E: Improvements & Trade-offs
        ("Q21: What are the main bottlenecks in your current system?", 
         "1. Sequential Graph Nodes: Agents run one after another rather than in parallel.\n2. Synchronous DB calls: `database.py` uses synchronous `psycopg.connect()` instead of an async connection pool (`AsyncConnectionPool`).\n3. External API Latency: RailKit and Tavily network lookups account for 80% of total response time."),

        ("Q22: How would you scale MoodyTrips to 100,000 active users?", 
         "1. Add Redis Caching: Cache common queries (e.g. weather in Goa or train schedules between Delhi and Jaipur) in Redis with 1-hour TTL.\n2. Async DB Connection Pooling: Use `psycopg_pool.AsyncConnectionPool` to handle concurrent FastAPI workers.\n3. Server-Sent Events (SSE) / WebSockets: Stream agent progress live to the browser so users see real-time updates as each agent completes."),

        ("Q23: How would you add user feedback or dynamic agent routing?", 
         "We can introduce a `Supervisor Agent` at the start of the LangGraph graph. The Supervisor analyzes the user query and dynamically routes to specific agents using conditional edges (`add_conditional_edges`). For example, if a user asks 'What's the weather in Tokyo?', the Supervisor routes directly to `weather_agent` -> `final_agent`, skipping flight and hotel agents completely!"),

        ("Q24: What metrics would you track to monitor system health in production?", 
         "1. Graph Latency per node (measuring execution time of Flight vs Hotel vs LLM).\n2. API Success/Failure Rate for Tavily, AviationStack, and RailKit.\n3. Token usage and costs per user request via `llm_calls` state tracking.\n4. Database checkpointer storage growth rate."),

        ("Q25: What was the biggest technical challenge you faced while building MoodyTrips?", 
         "The biggest challenge was state synchronization and error handling across heterogeneous external APIs. AviationStack uses IATA city codes, RailKit requires station codes, and Tavily expects natural language string queries. Resolving user input into structured parameters reliably using regex parsers and fallback search mechanisms without breaking the LangGraph checkpoint state required careful engineering and modular tool isolation.")
    ]

    for q, a in qa_list:
        add_h2(q)
        add_p(a)

    # ==========================================
    # SECTION 7: FUTURE IMPROVEMENTS & ARCHITECTURAL ROADMAP
    # ==========================================
    add_h1("SECTION 7: System Bottlenecks & Future Improvements")
    
    add_h2("7.1 Parallel Node Execution (LangGraph Fan-Out)")
    add_p("Transform sequential execution into parallel branch execution using LangGraph native branching:")
    add_code_block(
"""# Current Sequential Execution:
graph.add_edge(START, "flight_agent")
graph.add_edge("flight_agent", "hotel_agent")
graph.add_edge("hotel_agent", "ground_transport_agent")
graph.add_edge("ground_transport_agent", "weather_agent")
graph.add_edge("weather_agent", "budget_agent")

# Proposed Parallel Execution (Fan-Out / Fan-In):
graph.add_edge(START, "flight_agent")
graph.add_edge(START, "hotel_agent")
graph.add_edge(START, "ground_transport_agent")
graph.add_edge(START, "weather_agent")

# Join node (Fan-In):
graph.add_edge(["flight_agent", "hotel_agent", "ground_transport_agent", "weather_agent"], "budget_agent")"""
    )
    
    add_h2("7.2 System Roadmap Matrix")
    add_bullet("Store frequent train schedules, hotel search results, and weather forecasts in Redis with TTL to reduce external API costs and latency.", "1. Redis Distributed Caching: ")
    add_bullet("Implement Server-Sent Events (`EventSource`) in FastAPI so the frontend displays live loading indicators as each agent node completes.", "2. Real-Time Streaming UI: ")
    add_bullet("Integrate Pinecone or Qdrant vector database to store user travel preferences, past itineraries, and feedback for personalized recommendations.", "3. Vector DB Memory (RAG): ")
    add_bullet("Add Google OAuth2 login flow alongside custom email/password authentication.", "4. OAuth2 Social Login: ")

    doc.save("MoodyTrips_Fresher_Interview_Guide.docx")
    print("Successfully generated MoodyTrips_Fresher_Interview_Guide.docx")

if __name__ == "__main__":
    build_interview_doc()
